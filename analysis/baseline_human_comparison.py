"""Analyze blinded human ratings for LLM-vs-template baseline scenarios."""

from __future__ import annotations

import argparse
import csv
import json
import math
import random
from collections import defaultdict
from pathlib import Path
from typing import Any

import numpy as np
from scipy import stats


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "annotation" / "raw" / "baseline_comparison" / "baseline_comparison_packet_COMPLETED_independent_ratings.xlsx"
BLIND_KEY = ROOT / "annotation" / "study_packet_baseline_comparison" / "baseline_comparison_blind_key_DO_NOT_SEND.csv"
RESULTS = ROOT / "results" / "baseline_human_comparison.json"
AGGREGATE_CSV = ROOT / "results" / "baseline_human_comparison_summary.csv"
TABLE = ROOT / "paper" / "ieee_access_overleaf" / "table_baseline_human_comparison.tex"

REVIEWER_SHEETS = [
    "Security_Engineer_1",
    "Security_Engineer_2",
    "CS_Lecturer",
    "CyberMSc_Student_1",
    "CyberMSc_Student_2",
]

DIMENSIONS = [
    ("attck_alignment_1_5", "ATT\\&CK alignment"),
    ("stage_sequence_1_5", "Stage sequence"),
    ("narrative_detail_1_5", "Narrative detail"),
    ("overall_realism_1_5", "Overall realism"),
]


def load_workbook_rows(path: Path) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read completed .xlsx rating workbooks") from exc
    wb = load_workbook(path, data_only=True)
    sheets: dict[str, list[dict[str, Any]]] = {}
    for name in REVIEWER_SHEETS:
        if name not in wb.sheetnames:
            raise ValueError(f"missing reviewer sheet: {name}")
        ws = wb[name]
        headers = [str(ws.cell(1, col).value or "").strip() for col in range(1, ws.max_column + 1)]
        rows = []
        for r in range(2, ws.max_row + 1):
            row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
            if row.get("blind_id"):
                rows.append(row)
        sheets[name] = rows
    return sheets


def load_key(path: Path) -> dict[str, dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return {row["blind_id"]: row for row in csv.DictReader(f)}


def parse_rating(value: Any, *, sheet: str, blind_id: str, column: str) -> int:
    try:
        rating = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{sheet}:{blind_id}:{column} is not an integer rating: {value!r}") from exc
    if not 1 <= rating <= 5:
        raise ValueError(f"{sheet}:{blind_id}:{column} outside 1--5: {rating}")
    return rating


def validate_and_long_rows(sheets: dict[str, list[dict[str, Any]]], key: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    expected_ids = set(key)
    long_rows: list[dict[str, Any]] = []
    for sheet, rows in sheets.items():
        ids = {str(row.get("blind_id", "")).strip() for row in rows}
        missing = sorted(expected_ids - ids)
        extra = sorted(ids - expected_ids)
        if missing or extra:
            raise ValueError(f"{sheet} ID mismatch: missing={missing[:5]} extra={extra[:5]}")
        if len(rows) != len(expected_ids):
            raise ValueError(f"{sheet} row count {len(rows)} != expected {len(expected_ids)}")
        for row in rows:
            blind_id = str(row["blind_id"]).strip()
            for column, _label in DIMENSIONS:
                long_rows.append(
                    {
                        "reviewer": sheet,
                        "blind_id": blind_id,
                        "pair_id": key[blind_id]["pair_id"],
                        "source": key[blind_id]["source"],
                        "dimension": column,
                        "rating": parse_rating(row.get(column), sheet=sheet, blind_id=blind_id, column=column),
                    }
                )
    return long_rows


def mean(values: list[float]) -> float:
    return float(np.mean(values)) if values else float("nan")


def bootstrap_ci(values: list[float], seed: int, n_boot: int = 5000) -> list[float]:
    if not values:
        return [float("nan"), float("nan")]
    rng = random.Random(seed)
    vals = []
    n = len(values)
    for _ in range(n_boot):
        sample = [values[rng.randrange(n)] for _ in range(n)]
        vals.append(mean(sample))
    lo, hi = np.percentile(vals, [2.5, 97.5])
    return [round(float(lo), 6), round(float(hi), 6)]


def summarize_dimension(long_rows: list[dict[str, Any]], dimension: str, seed: int) -> dict[str, Any]:
    scenario_values: dict[tuple[str, str, str], list[int]] = defaultdict(list)
    for row in long_rows:
        if row["dimension"] == dimension:
            scenario_values[(row["pair_id"], row["blind_id"], row["source"])].append(row["rating"])

    pair_values: dict[str, dict[str, float]] = defaultdict(dict)
    for (pair_id, _blind_id, source), vals in scenario_values.items():
        pair_values[pair_id][source] = mean(vals)

    diffs = []
    llm_vals = []
    template_vals = []
    for pair_id, sources in sorted(pair_values.items()):
        if "LLM" not in sources or "TEMPLATE" not in sources:
            raise ValueError(f"pair {pair_id} missing source rating: {sources}")
        llm_vals.append(sources["LLM"])
        template_vals.append(sources["TEMPLATE"])
        diffs.append(sources["LLM"] - sources["TEMPLATE"])

    wilcoxon = stats.wilcoxon(diffs, alternative="greater", zero_method="zsplit")
    ttest = stats.ttest_rel(llm_vals, template_vals, alternative="greater")
    return {
        "pairs": len(diffs),
        "llm_mean": round(mean(llm_vals), 6),
        "template_mean": round(mean(template_vals), 6),
        "mean_difference": round(mean(diffs), 6),
        "mean_difference_95ci": bootstrap_ci(diffs, seed),
        "median_difference": round(float(np.median(diffs)), 6),
        "pairs_llm_greater": sum(1 for d in diffs if d > 0),
        "pairs_equal": sum(1 for d in diffs if d == 0),
        "pairs_template_greater": sum(1 for d in diffs if d < 0),
        "wilcoxon_p_greater": float(wilcoxon.pvalue),
        "paired_t_p_greater": float(ttest.pvalue),
        "diffs": [round(float(d), 6) for d in diffs],
    }


def write_summary_csv(long_rows: list[dict[str, Any]], key: dict[str, dict[str, str]]) -> None:
    by_scenario_dim: dict[tuple[str, str], list[int]] = defaultdict(list)
    for row in long_rows:
        by_scenario_dim[(row["blind_id"], row["dimension"])].append(row["rating"])
    rows = []
    for blind_id, meta in sorted(key.items(), key=lambda item: (item[1]["pair_id"], item[1]["source"])):
        out = {**meta}
        for column, _label in DIMENSIONS:
            out[f"mean_{column}"] = round(mean(by_scenario_dim[(blind_id, column)]), 6)
        rows.append(out)
    AGGREGATE_CSV.parent.mkdir(parents=True, exist_ok=True)
    with AGGREGATE_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def fmt(value: float, digits: int = 2) -> str:
    return f"{value:.{digits}f}"


def write_table(result: dict[str, Any]) -> None:
    lines = [
        "% Auto-generated by analysis/baseline_human_comparison.py",
        "\\begin{table*}[t]",
        "\\centering",
        "\\caption{Blinded human comparison of matched LLM and ATT\\&CK-constrained template scenarios. Positive differences favour the LLM scenario in each pair.}",
        "\\label{tab:baseline-human-comparison}",
        "\\small",
        "\\begin{tabular}{lrrrrr}",
        "\\toprule",
        "Dimension & LLM mean & Template mean & Mean diff. [95\\% CI] & LLM $>$ template & Wilcoxon $p$ \\\\",
        "\\midrule",
    ]
    for column, label in DIMENSIONS:
        row = result["dimensions"][column]
        lo, hi = row["mean_difference_95ci"]
        lines.append(
            f"{label} & {fmt(row['llm_mean'])} & {fmt(row['template_mean'])} & "
            f"{fmt(row['mean_difference'])} [{fmt(lo)}, {fmt(hi)}] & "
            f"{row['pairs_llm_greater']}/{row['pairs']} & {row['wilcoxon_p_greater']:.3g} \\\\"
        )
    lines.extend(["\\bottomrule", "\\end{tabular}", "\\end{table*}", ""])
    TABLE.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--blind-key", type=Path, default=BLIND_KEY)
    parser.add_argument("--seed", type=int, default=20260617)
    args = parser.parse_args()

    key = load_key(args.blind_key)
    sheets = load_workbook_rows(args.workbook)
    long_rows = validate_and_long_rows(sheets, key)
    dimensions = {
        column: summarize_dimension(long_rows, column, args.seed + idx)
        for idx, (column, _label) in enumerate(DIMENSIONS)
    }
    result = {
        "schema_version": 1,
        "status": "complete",
        "source_workbook": str(args.workbook),
        "scenario_count": len(key),
        "matched_pairs": len({row["pair_id"] for row in key.values()}),
        "reviewer_count": len(REVIEWER_SHEETS),
        "reviewer_profile": {
            "roles": {
                "security_tester": 2,
                "computer_science_lecturer": 1,
                "cybersecurity_msc_student": 2,
            },
            "professional_experience_years": "2-6",
            "participation": "voluntary",
            "author_status": "No reviewer was an author of the paper.",
        },
        "dimensions": dimensions,
        "validation": {
            "reviewer_sheets": REVIEWER_SHEETS,
            "expected_blind_ids": len(key),
            "ratings_per_scenario_per_dimension": len(REVIEWER_SHEETS),
            "rating_range": [1, 5],
            "note": "Consensus Rating_Template sheet is ignored; only independent reviewer sheets are analyzed.",
        },
    }
    RESULTS.parent.mkdir(parents=True, exist_ok=True)
    RESULTS.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    write_summary_csv(long_rows, key)
    write_table(result)
    print(f"validated {len(REVIEWER_SHEETS)} reviewer sheets, {len(key)} scenarios")
    print(f"wrote {RESULTS}")
    print(f"wrote {AGGREGATE_CSV}")
    print(f"wrote {TABLE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
